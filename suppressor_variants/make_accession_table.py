#!/usr/bin/env python3
"""Supplementary table with the NCBI accessions of the whole-genome sequencing data.

Accessions come from the SRA run metadata of BioProject PRJNA1418127 (checked 2026-09-29).
Writes Supplementary_Table_WGS_accessions.xlsx and .tsv.
"""
import os
import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
BIOPROJECT = "PRJNA1418127"

# strain, description, parental strain, BioSample, SRA sample, SRA experiment, SRA run, bases
ROWS = [
    ("WT", "Wild-type (BY4741)", "-", "SAMN03020231", "SRS697355", "SRX32029693", "SRR37083327", 571234717),
    ("tif51A-1", "tif51A-1 (BY4741 background)", "-", "SAMN55031537", "SRS27969596", "SRX32029694", "SRR37083326", 639036037),
    ("tif51A-3", "tif51A-3 (BY4741 background)", "-", "SAMN55031538", "SRS27969598", "SRX32029696", "SRR37083324", 529339535),
    ("sup.1", "Suppressor of tif51A-1", "tif51A-1", "SAMN55031539", "SRS27969599", "SRX32029697", "SRR37083323", 439056297),
    ("sup.2", "Suppressor of tif51A-1", "tif51A-1", "SAMN55031540", "SRS27969600", "SRX32029698", "SRR37083322", 626780695),
    ("sup.22", "Suppressor of tif51A-1", "tif51A-1", "SAMN55031541", "SRS27969601", "SRX32029699", "SRR37083321", 608183324),
    ("sup.23", "Suppressor of tif51A-1", "tif51A-1", "SAMN55031542", "SRS27969602", "SRX32029700", "SRR37083320", 568165323),
    ("sup.11", "Suppressor of tif51A-3", "tif51A-3", "SAMN55031543", "SRS27969603", "SRX32029701", "SRR37083319", 688946917),
    ("sup.15", "Suppressor of tif51A-3", "tif51A-3", "SAMN55031544", "SRS27969604", "SRX32029702", "SRR37083318", 489646015),
    ("sup.25", "Suppressor of tif51A-3", "tif51A-3", "SAMN55031545", "SRS27969605", "SRX32029703", "SRR37083317", 544733853),
    ("sup.27", "Suppressor of tif51A-3", "tif51A-3", "SAMN55031546", "SRS27969597", "SRX32029695", "SRR37083325", 608921843),
]
HEADER = ["Strain", "Description", "Parental strain", "BioProject", "BioSample", "SRA run"]


def values(r):
    strain, desc, parent, biosample, srs, srx, srr, bases = r
    return [strain, desc, parent, BIOPROJECT, biosample, srr]


def main():
    out = os.path.join(HERE, "Supplementary_Table_WGS_accessions")
    with open(out + ".tsv", "w") as o:
        o.write("\t".join(HEADER) + "\n")
        for r in ROWS:
            o.write("\t".join(map(str, values(r))) + "\n")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "WGS accessions"
    ws.append(HEADER)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in ROWS:
        ws.append(values(r))
    ws.append([])
    ws.append([f"All sequencing runs belong to BioProject {BIOPROJECT}. The wild-type run (SRR37083327) is linked to "
               "the existing BioSample SAMN03020231 (BY4741); all other BioSamples were created for this study."])
    widths = [10, 30, 14, 14, 15, 14]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "B2"
    wb.save(out + ".xlsx")
    print(f"written: {out}.xlsx\n         {out}.tsv")


if __name__ == "__main__":
    main()
