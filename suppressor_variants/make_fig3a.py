#!/usr/bin/env python3
"""Redraw Figure 3a from Supplementary_Table_WGS_variants.xlsx (find_suppressor_variants.py).

A gene/strain cell is filled when the strain carries a variant marked "In Fig. 3a":
suppressor-specific (all three controls genotyped as reference with >= 10 reads and <= 2 ALT
reads) and protein-altering according to SnpEff.
SNVs go to the SNPs panel; insertions and deletions to the INDELs panel.
"""
import os
from collections import defaultdict
import openpyxl
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle

HERE = os.path.dirname(os.path.abspath(__file__))
STRAINS = ["sup.1", "sup.2", "sup.22", "sup.23", "sup.11", "sup.15", "sup.25", "sup.27"]
LINEAGE = {"tif51A-1": STRAINS[:4], "tif51A-3": STRAINS[4:]}
LINEAGE_COLOR = {"tif51A-1": "#D9731A", "tif51A-3": "#A21CAF"}
FILL = {"ROX1": "#E6337F", "MOT3": "#3F7FE0"}
OTHER, EMPTY = "#F2A65A", "#E3E3E3"
FIRST = ["ROX1", "MOT3"]  # candidate genes listed first

ws = openpyxl.load_workbook(os.path.join(HERE, "Supplementary_Table_WGS_variants.xlsx"))["Variants"]
hdr = [c.value for c in ws[1]]
col = {h: i for i, h in enumerate(hdr)}

cells = {"SNPs": defaultdict(dict), "INDELs": defaultdict(dict)}
for r in ws.iter_rows(min_row=2, values_only=True):
    if r[col["In Fig. 3a"]] != "yes":
        continue
    gene = r[col["Gene (SGD)"]]
    panel = "SNPs" if r[col["Type"]] == "SNV" else "INDELs"
    for s in r[col["Suppressor(s)"]].split(", "):
        cells[panel][gene][s] = False  # False = no asterisk (kept for the drawing code below)


def order(genes):
    return [g for g in FIRST if g in genes] + sorted(g for g in genes if g not in FIRST)

plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 9, "pdf.fonttype": 42, "svg.fonttype": "none"})
fig, axes = plt.subplots(1, 2, figsize=(7.2, 2.2), gridspec_kw={"wspace": 0.45})
gap = 0.06
for ax, panel in zip(axes, ["SNPs", "INDELs"]):
    genes = order(cells[panel])
    for yi, g in enumerate(genes):
        for xi, s in enumerate(STRAINS):
            low = cells[panel][g].get(s)
            color = EMPTY if low is None else FILL.get(g, OTHER)
            ax.add_patch(Rectangle((xi + gap / 2, yi + gap / 2), 1 - gap, 1 - gap, color=color, lw=0))
            if low:
                ax.text(xi + 0.5, yi + 0.55, "*", ha="center", va="center", color="white",
                        fontsize=11, fontweight="bold")
    ax.set_xlim(0, len(STRAINS)); ax.set_ylim(len(genes), 0)
    ax.set_yticks([i + 0.5 for i in range(len(genes))])
    ax.set_yticklabels(genes, fontstyle="italic")
    for lab, g in zip(ax.get_yticklabels(), genes):
        if g in FILL:
            lab.set_color(FILL[g]); lab.set_fontweight("bold")
    ax.set_xticks([i + 0.5 for i in range(len(STRAINS))])
    ax.set_xticklabels([s.split(".")[1] for s in STRAINS])
    for lab, s in zip(ax.get_xticklabels(), STRAINS):
        lab.set_color(LINEAGE_COLOR["tif51A-1" if s in LINEAGE["tif51A-1"] else "tif51A-3"])
    ax.tick_params(length=0, pad=3)
    for sp in ax.spines.values():
        sp.set_visible(False)
    ax.set_title(panel, fontsize=10)
    ax.set_aspect("equal"); ax.set_anchor("N")
    # lineage brackets below the strain numbers
    for li, (lin, members) in enumerate(LINEAGE.items()):
        x0, x1 = STRAINS.index(members[0]) + 0.1, STRAINS.index(members[-1]) + 0.9
        y = len(genes) + 0.95 + (0 if len(genes) > 3 else 0)
        ax.plot([x0, x1], [y, y], color="#555555", lw=0.8, clip_on=False)
        ax.text((x0 + x1) / 2, y + 0.25, lin, ha="center", va="top", fontstyle="italic", clip_on=False)

fig.text(0.02, 0.97, "a", fontsize=16, fontweight="bold", va="top")
for ext in ("pdf", "svg", "png"):
    fig.savefig(os.path.join(HERE, f"Figure3a_corrected.{ext}"), dpi=600 if ext == "png" else None,
                bbox_inches="tight")
for panel in cells:
    for g in order(cells[panel]):
        print(panel, g, {s: ("low-DP*" if v else "ok") for s, v in cells[panel][g].items()})
