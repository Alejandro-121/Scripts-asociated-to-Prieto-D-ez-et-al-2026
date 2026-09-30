#!/usr/bin/env python3
"""Identify suppressor-specific variants in the joint-genotyped, SnpEff-annotated VCF.

Input: cohort.filtered.ann.vcf.gz (run_variant_calling.sh on BQSR BAMs + annotate_snpeff.sh),
genotyped jointly in all 11 strains from per-base (BP_RESOLUTION) GVCFs, so every strain
has a genotype and its own read depth at every variant site.

Candidates: every ALT allele genotyped in >= 1 suppressor and in none of the control strains
(WT, tif51A-1 = 2-1, tif51A-3 = 2-3). Each candidate is classified in this order:
  background      - a control has > MAX_CTRL_ALT reads supporting the ALT allele, or carries another
                    ALT allele at the site (includes controls left uncalled './.' whose reads carry it)
  not testable    - a control is not called as reference (GT './.') or has < MIN_CTRL_DP reads
  suppressor-specific - all three controls GT = 0 with >= MIN_CTRL_DP reads and <= MAX_CTRL_ALT ALT reads
Suppressor-specific variants are kept when the evidence in the suppressor is sufficient: the site passes
the GATK hard filters, the ALT allele is in >= MIN_AF of the reads of every carrier (haploid, clonal
strains) and at least one carrier has >= MIN_CARRIER_DP reads. Only these are written to the table.
Experimental validation: alleles reconstructed by CRISPR-Cas9 in the parental strains (Fig. 4).
Only the VCF is used; the BAM files are not read.

Writes <out>.tsv and <out>.xlsx (sheets: Variants, Summary, Criteria).
"""
import argparse, gzip, os, re, sys
from collections import Counter, defaultdict
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))

CONTROLS = {"WT": "WT", "2-1": "tif51A-1", "2-3": "tif51A-3"}
SUPPRESSORS = ["sup.1", "sup.2", "sup.22", "sup.23", "sup.11", "sup.15", "sup.25", "sup.27"]
PARENT = {s: "tif51A-1" for s in SUPPRESSORS[:4]} | {s: "tif51A-3" for s in SUPPRESSORS[4:]}
PROTEIN_ALTERING = ("missense_variant", "frameshift_variant", "stop_gained", "stop_lost", "start_lost",
                    "inframe_insertion", "inframe_deletion", "disruptive_inframe", "conservative_inframe",
                    "splice_acceptor", "splice_donor")
STATUS_ORDER = ["suppressor-specific", "not testable", "background"]
FILLS = {"suppressor-specific": "E2EFDA", "not testable": "EDEDED", "background": "F8CBAD"}
# (chrom, pos, REF, ALT tras recortar) -> validación experimental descrita en el manuscrito (Fig. 4)
CRISPR = "Reconstructed by CRISPR-Cas9 in the parental {} strain (Sanger-confirmed); reproduces suppression and TIF51B derepression (Fig. 4)"
VALIDATED = {
    ("chrXVI", 679850, "G", "T"): CRISPR.format("tif51A-1 ROX1-13myc") + " [SNP1]",
    ("chrXVI", 679910, "T", "C"): CRISPR.format("tif51A-1 ROX1-13myc") + " [SNP2]",
    ("chrXVI", 679884, "G", "T"): CRISPR.format("tif51A-3 ROX1-13myc") + " [SNP3]",
    ("chrXIII", 409850, "G", "GCCCCGGCCCCCGGTCCA"): CRISPR.format("tif51A-3") + " [INS1]",
}


def parse_args():
    p = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    p.add_argument("-v", "--vcf", default=os.path.join(HERE, "vc_final/vcf/cohort.filtered.ann.vcf.gz"))
    p.add_argument("-o", "--out", default=os.path.join(HERE, "Supplementary_Table_WGS_variants"),
                   help="output prefix (.tsv and .xlsx are added)")
    p.add_argument("--min-ctrl-dp", type=int, default=10, help="minimum reads in every control (default 10)")
    p.add_argument("--max-ctrl-alt", type=int, default=2, help="maximum ALT reads in any control (default 2)")
    p.add_argument("--min-af", type=float, default=0.8,
                   help="minimum ALT fraction in every carrier (default 0.8)")
    p.add_argument("--min-carrier-dp", type=int, default=10,
                   help="minimum reads in at least one carrier (default 10)")
    p.add_argument("--gff", default=os.path.join(HERE, "saccharomyces_cerevisiae.gff.gz"),
                   help="current SGD GFF3, only used for present-day gene names")
    return p.parse_args()


def trim(ref, alt):
    """Remove the suffix shared by REF and ALT (multi-allelic records)."""
    while len(ref) > 1 and len(alt) > 1 and ref[-1] == alt[-1]:
        ref, alt = ref[:-1], alt[:-1]
    return ref, alt


def parse_info(field):
    info = {}
    for kv in field.split(";"):
        k, _, v = kv.partition("=")
        info[k] = v if _ else True
    return info


def read_vcf(path):
    filters, samples, records = {}, [], []
    with gzip.open(path, "rt") as fh:
        for line in fh:
            if line.startswith("##FILTER=<ID="):
                m = re.match(r'##FILTER=<ID=([^,]+),Description="([^"]*)"', line)
                filters[m.group(1)] = m.group(2)
                continue
            if line.startswith("##"):
                continue
            f = line.rstrip("\n").split("\t")
            if line.startswith("#"):
                samples = f[9:]
                continue
            fmt = f[8].split(":")
            gts = {s: dict(zip(fmt, v.split(":"))) for s, v in zip(samples, f[9:])}
            records.append((f[0], int(f[1]), f[3], f[4].split(","), f[5], f[6], parse_info(f[7]), gts))
    return filters, samples, records


def allele_reads(g, k, n_alts):
    """(ALT reads for allele k, total informative reads) from AD."""
    ad = g.get("AD", ".")
    if ad in (".", ""):
        return 0, 0
    v = [int(x) if x != "." else 0 for x in ad.split(",")]
    v += [0] * (n_alts + 1 - len(v))
    return v[k], sum(v)


IMPACT_RANK = {"HIGH": 0, "MODERATE": 1, "LOW": 2, "MODIFIER": 3}
EFFECT_RANK = ["splice", "5_prime_UTR", "3_prime_UTR", "intron", "non_coding", "upstream", "downstream", "intergenic"]


def snpeff(info, alt, dubious=frozenset()):
    """Most relevant SnpEff annotation for this ALT allele.

    SnpEff lists every gene within 5 kb; in the compact yeast genome that is several genes per
    variant, so the annotation is chosen by impact, then by effect (coding > splice > UTR > intron >
    up/downstream > intergenic), then by distance to the gene. Dubious ORFs go last, so that a variant
    in TIF51A (HYP2) is not reported as a stop codon in the overlapping dubious ORF YEL034C-A.
    """
    best, best_key = None, None
    for ann in info.get("ANN", "").split(","):
        a = ann.split("|")
        if len(a) < 15 or a[0] != alt:
            continue
        eff = next((i for i, e in enumerate(EFFECT_RANK) if e in a[1]), -1)
        dist = int(a[14]) if a[14].isdigit() else 0
        key = (a[4] in dubious, IMPACT_RANK.get(a[2], 4), eff, dist)
        if best_key is None or key < best_key:
            best, best_key = a, key
    if best is None:
        return dict(effect="", impact="", gene="", gene_id="", biotype="", hgvs_c="", hgvs_p="", distance="")
    return dict(effect=best[1], impact=best[2], gene=best[3], gene_id=best[4], biotype=best[7],
                hgvs_c=best[9], hgvs_p=best[10], distance=best[14])


def current_names(gff):
    """Systematic ID -> current SGD standard name (the R64-1-1 GFF of 2011 lacks later names),
    and the set of ORFs classified as Dubious."""
    names, dubious = {}, set()
    if not os.path.exists(gff):
        return names, dubious
    with gzip.open(gff, "rt") as fh:
        for line in fh:
            if line.startswith("##FASTA"):
                break
            f = line.split("\t")
            if len(f) == 9 and f[2] == "gene":
                attr = dict(kv.split("=", 1) for kv in f[8].strip().split(";") if "=" in kv)
                if attr.get("gene"):
                    names[attr["ID"]] = attr["gene"]
                if attr.get("orf_classification") == "Dubious":
                    dubious.add(attr["ID"])
    return names, dubious


def fmt_num(x):
    try:
        return round(float(x), 2)
    except (TypeError, ValueError):
        return ""


def classify(ctrl, a):
    if any(c["alt"] > a.max_ctrl_alt or c["gt"] not in ("0", ".") for c in ctrl.values()):
        return "background"
    if any(c["gt"] != "0" or c["dp"] < a.min_ctrl_dp for c in ctrl.values()):
        return "not testable"
    return "suppressor-specific"


def evidence(x, a):
    """Motivo por el que la evidencia en el supresor es insuficiente ('' si es suficiente)."""
    if x["filt"] != "PASS":
        return "fails GATK hard filters"
    if any(x["car"][s]["dp"] == 0 or x["car"][s]["alt"] / x["car"][s]["dp"] < a.min_af for s in x["carriers"]):
        return f"ALT frequency < {a.min_af} in a carrier"
    if max(x["car"][s]["dp"] for s in x["carriers"]) < a.min_carrier_dp:
        return f"< {a.min_carrier_dp} reads in every carrier"
    return ""


def main():
    a = parse_args()
    filters, samples, records = read_vcf(a.vcf)
    names, dubious = current_names(a.gff)
    missing = [s for s in list(CONTROLS) + SUPPRESSORS if s not in samples]
    if missing:
        sys.exit(f"ERROR: faltan muestras en el VCF: {missing}")

    rows = []
    for chrom, pos, ref, alts, qual, filt, info, gts in records:
        for k, alt in enumerate(alts, 1):
            if alt == "*":
                continue
            carriers = [s for s in SUPPRESSORS if gts[s].get("GT") == str(k)]
            if not carriers or any(gts[c].get("GT") == str(k) for c in CONTROLS):
                continue
            r, al = trim(ref, alt)
            ctrl = {}
            for c in CONTROLS:
                alt_n, dp = allele_reads(gts[c], k, len(alts))
                ctrl[c] = dict(gt=gts[c].get("GT", "."), alt=alt_n, dp=dp)
            car = {}
            for s in carriers:
                alt_n, dp = allele_reads(gts[s], k, len(alts))
                car[s] = dict(alt=alt_n, dp=dp)
            ann = snpeff(info, alt, dubious)
            ann["current"] = names.get(ann["gene_id"], "")
            rows.append(dict(
                chrom=chrom, pos=pos, ref=r, alt=al, vtype="SNV" if len(r) == len(al) == 1 else
                ("insertion" if len(al) > len(r) else "deletion"),
                carriers=carriers, car=car, ctrl=ctrl, qual=fmt_num(qual), filt=filt,
                qd=fmt_num(info.get("QD")), fs=fmt_num(info.get("FS")), sor=fmt_num(info.get("SOR")),
                mq=fmt_num(info.get("MQ")), **ann,
                protein=any(e in ann["effect"] for e in PROTEIN_ALTERING),
                status=classify(ctrl, a)))
            rows[-1]["weak"] = evidence(rows[-1], a)

    for x in rows:
        x["validation"] = VALIDATED.get((x["chrom"], x["pos"], x["ref"], x["alt"]), "Not tested")
        x["fig"] = x["status"] == "suppressor-specific" and x["protein"]

    # la tabla solo lleva las específicas con evidencia suficiente; el resto queda como recuento en Summary
    tally = Counter(x["status"] for x in rows)
    specific = [x for x in rows if x["status"] == "suppressor-specific"]
    tally.update(x["weak"] for x in specific if x["weak"])
    rows = [x for x in specific if not x["weak"]]
    rows.sort(key=lambda x: (not x["fig"], x["current"] or x["gene"] or "~", x["chrom"], x["pos"]))
    write_outputs(a, rows, filters, tally)
    report(rows, tally)


HEADER = ["In Fig. 3a", "Experimental validation", "Gene (SGD)", "Systematic ID", "SnpEff effect", "SnpEff impact",
          "SnpEff distance to gene (bp)", "SnpEff coding DNA change (HGVS c.)",
          "SnpEff protein change (HGVS p.)", "Chromosome", "Position (R64-1-1)", "REF", "ALT", "Type", "Suppressor(s)",
          "Parental background", "Suppressor ALT/DP (VCF)", "Suppressor allele frequency",
          "WT GT", "WT ALT/DP", "tif51A-1 GT", "tif51A-1 ALT/DP", "tif51A-3 GT", "tif51A-3 ALT/DP",
          "QUAL", "QD", "FS", "SOR", "MQ", "GATK FILTER"]


def row_values(x):
    af = [round(x["car"][s]["alt"] / x["car"][s]["dp"], 2) if x["car"][s]["dp"] else "" for s in x["carriers"]]
    return ["yes" if x["fig"] else "", x["validation"], x["current"] or x["gene"], x["gene_id"], x["effect"], x["impact"],
            x["distance"], x["hgvs_c"], x["hgvs_p"],
            x["chrom"], x["pos"], x["ref"], x["alt"], x["vtype"], ", ".join(x["carriers"]),
            ", ".join(sorted({PARENT[s] for s in x["carriers"]})),
            "; ".join(f"{s}: {x['car'][s]['alt']}/{x['car'][s]['dp']}" for s in x["carriers"]),
            "; ".join(map(str, af)),
            *[v for c in CONTROLS for v in (x["ctrl"][c]["gt"], f"{x['ctrl'][c]['alt']}/{x['ctrl'][c]['dp']}")],
            x["qual"], x["qd"], x["fs"], x["sor"], x["mq"], x["filt"]]


def write_outputs(a, rows, filters, tally):
    with open(a.out + ".tsv", "w") as o:
        o.write("\t".join(HEADER) + "\n")
        for x in rows:
            o.write("\t".join(map(str, row_values(x))) + "\n")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Variants"
    ws.append(HEADER)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for x in rows:
        ws.append(row_values(x))
        i = ws.max_row
        if x["fig"]:
            for j in (1, 3):
                ws.cell(row=i, column=j).font = Font(bold=True)
                ws.cell(row=i, column=j).fill = PatternFill("solid", fgColor=FILLS["suppressor-specific"])
    widths = {1: 9, 2: 40, 3: 10, 4: 12, 5: 28, 6: 9, 8: 22, 9: 18, 12: 10, 13: 10, 15: 16, 17: 22}
    for j in range(1, len(HEADER) + 1):
        ws.column_dimensions[get_column_letter(j)].width = widths.get(j, 12)
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions

    # Summary: cómo se llegó a las específicas y mutaciones que alteran la proteína por supresor
    ws2 = wb.create_sheet("Summary")
    ws2.append(["Candidate alleles (in >= 1 suppressor, not genotyped in any control)", sum(tally[s] for s in STATUS_ORDER)])
    ws2.append(["  excluded: background (variant present in a control)", tally["background"]])
    ws2.append(["  excluded: not testable (control no-call or < 10 reads)", tally["not testable"]])
    n_spec = tally["suppressor-specific"]
    ws2.append(["Suppressor-specific (controls genotyped as reference, >= 10 reads, <= 2 ALT reads)", n_spec])
    for why in (k for k in tally if k not in STATUS_ORDER):
        ws2.append([f"  excluded: insufficient evidence in the suppressor ({why})", tally[why]])
    ws2.append(["Suppressor-specific variants with sufficient evidence (this table)", len(rows)])
    ws2.append(["  of which protein-altering (Fig. 3a)", sum(x["fig"] for x in rows)])
    ws2.append([])
    ws2.append(["Suppressor", "Parental background", "Suppressor-specific protein-altering variants (gene: protein change)"])
    hdr_row = ws2.max_row
    per = defaultdict(list)
    for x in rows:
        if x["fig"]:
            for s in x["carriers"]:
                per[s].append(f"{x['current'] or x['gene']}: {x['hgvs_p'] or x['hgvs_c']}")
    for s in SUPPRESSORS:
        ws2.append([s, PARENT[s], "; ".join(per[s]) or "-"])
    for r in (ws2.max_row - len(SUPPRESSORS) - 3, hdr_row):
        for c in ws2[r]:
            c.font = Font(bold=True)
    ws2.column_dimensions["A"].width = 80
    ws2.column_dimensions["B"].width = 20
    ws2.column_dimensions["C"].width = 90

    # Criteria
    ws3 = wb.create_sheet("Criteria")
    lines = [
        "Reads: FastQC/MultiQC; first 15 bp trimmed with Trim Galore. Alignment: bwa-mem2 to S. cerevisiae S288C "
        "R64-1-1 (nuclear chromosomes); duplicates marked with GATK MarkDuplicates.",
        "Base quality score recalibration: GATK BaseRecalibrator/ApplyBQSR using the PASS variants of an initial "
        "joint call as known sites (bootstrapping); evaluated with AnalyzeCovariates.",
        "Variant calling: GATK HaplotypeCaller (-ploidy 1, -ERC BP_RESOLUTION) per strain, CombineGVCFs and "
        "GenotypeGVCFs (-ploidy 1) jointly for all 11 strains, so every strain has a genotype and read depth at "
        "every variant site.",
        "Hard filters (GATK VariantFiltration; failing sites are flagged in FILTER): "
        + "; ".join(f"{k}: {v}" for k, v in filters.items() if k not in ("LowQual", "LowDP")) + ".",
        "Annotation: SnpEff with a custom database built from the R64-1-1 reference and the SGD R64-1-1 GFF3 "
        "(CDS and protein sequences checked against SGD orf_coding_all / orf_trans_all: 0 errors).",
        "Candidates: ALT alleles genotyped in >= 1 suppressor and in none of the controls (WT, tif51A-1, tif51A-3). "
        "Only suppressor-specific candidates are listed in this table; the number of excluded candidates is given in "
        "the Summary sheet.",
        f"background: a control has > {a.max_ctrl_alt} reads supporting the ALT allele (also when the control genotype "
        "is './.') or is genotyped as another ALT allele at the site.",
        f"not testable: a control is not confidently genotyped as reference (GT './.', i.e. GQ = 0) or has "
        f"< {a.min_ctrl_dp} reads at the site.",
        f"suppressor-specific: all three controls genotyped as reference (GT = 0) with >= {a.min_ctrl_dp} reads and "
        f"<= {a.max_ctrl_alt} ALT-supporting reads.",
        f"Sufficient evidence in the suppressor (required for this table): site passes the GATK hard filters, ALT "
        f"allele in >= {a.min_af:.0%} of the reads of every carrier (strains are haploid and clonal) and >= "
        f"{a.min_carrier_dp} reads in at least one carrier.",
        "In Fig. 3a: listed in this table and protein-altering (missense, nonsense, start/stop lost, frameshift, "
        "in-frame indel, splice site).",
        "Experimental validation: ROX1 SNP1 (W53L), SNP2 (L73P), SNP3 (W64C) and MOT3 INS1 were reconstructed by "
        "CRISPR-Cas9 in the parental strains, confirmed by Sanger sequencing, and reproduced the suppression phenotype "
        "and TIF51B derepression (Fig. 4). Other variants were not tested experimentally.",
        "Allele frequency = ALT reads / (REF + ALT reads) from the per-sample AD field.",
    ]
    ws3.append(["Criteria and methods"])
    ws3["A1"].font = Font(bold=True)
    for l in lines:
        ws3.append([l])
        ws3.cell(row=ws3.max_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws3.column_dimensions["A"].width = 130
    wb.save(a.out + ".xlsx")
    print(f"escrito: {a.out}.xlsx\n         {a.out}.tsv")


def report(rows, tally):
    print("candidatas:", sum(tally[s] for s in STATUS_ORDER), "|", ", ".join(f"{k}: {v}" for k, v in tally.items()),
          "| en la tabla:", len(rows))
    print("\nEspecíficas de supresor que alteran la proteína:")
    for x in rows:
        if x["status"] == "suppressor-specific" and x["protein"]:
            val = "CRISPR" if x["validation"] != "Not tested" else ""
            print(f"  {x['current'] or x['gene']:<9} {x['hgvs_p'] or x['hgvs_c']:<18} {x['effect'][:38]:<38} "
                  f"{','.join(x['carriers']):<14} {val}")


if __name__ == "__main__":
    main()
